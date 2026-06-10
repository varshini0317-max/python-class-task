# TO CREATE A DOC
from docx import Document
from docx.shared import Inches

doc = Document()
doc.save("C:\\Users\\Hp\\Desktop\\Resume.docx")
print("Resume created successfully")

# FOR PHOTO
photo = input("Enter photo path: ")
doc.add_picture(photo, width=Inches(1.5), height=Inches(2))
print("Photo added successfully")

# FOR NAME AND PERSONAL INFO
name = input("Enter Your Name: ")
phone = input("Enter Your Phone Number: ")
if len(phone) == 10 and phone.isdigit():
    print("Valid Phone Number")
else:
    print("Invalid Phone Number")
    exit()

email = input("Enter Email Address: ")
if "@" not in email:
    print("Invalid Email Address")
    exit()

address = input("Enter Address: ")
languages = input("Enter languages you know (comma separated): ")
hobbies = input("Enter your hobbies (comma separated): ")

# FOR SKILLS AND EXPERIENCE
education = input("Education: ")
college = input("College/School: ")
experience = input("Experience: ")
skills = input("Skills (comma separated): ")
projects = input("Projects: ")

# PROFILE SUMMARY
intro_para = f"""
I am a motivated and dedicated individual with a strong academic background in {education} from {college}. 
I have hands-on experience in {experience}, backed by a solid set of technical and practical skills. 
My core competencies include {skills}. I have successfully completed projects such as {projects}, 
which have sharpened my problem-solving and development abilities. 
I am fluent in {languages} and pursue interests such as {hobbies} outside of my professional work.
"""

# FOR RATING - calculated before use
skill_count = len(skills.split(","))
if skill_count >= 10:
    skill_stars = "★★★★★"
elif skill_count >= 8:
    skill_stars = "★★★★☆"
elif skill_count >= 6:
    skill_stars = "★★★☆☆"
elif skill_count >= 4:
    skill_stars = "★★☆☆☆"
else:
    skill_stars = "★☆☆☆☆"

project_count = len(projects.split(","))
if project_count >= 5:
    project_stars = "★★★★★"
elif project_count == 4:
    project_stars = "★★★★☆"
elif project_count == 3:
    project_stars = "★★★☆☆"
elif project_count == 2:
    project_stars = "★★☆☆☆"
else:
    project_stars = "★☆☆☆☆"

language_count = len(languages.split(","))
if language_count >= 5:
    language_stars = "★★★★★"
elif language_count == 4:
    language_stars = "★★★★☆"
elif language_count == 3:
    language_stars = "★★★☆☆"
elif language_count == 2:
    language_stars = "★★☆☆☆"
else:
    language_stars = "★☆☆☆☆"

experience_count = len(experience.split(","))
if experience_count >= 5:
    experience_stars = "★★★★★"
elif experience_count == 4:
    experience_stars = "★★★★☆"
elif experience_count == 3:
    experience_stars = "★★★☆☆"
elif experience_count == 2:
    experience_stars = "★★☆☆☆"
else:
    experience_stars = "★☆☆☆☆"

# BUILD DOCUMENT
doc.add_heading(name, 0)
doc.add_heading("RESUME", 0)

doc.add_heading("Profile Summary", 1)
doc.add_paragraph(intro_para)

doc.add_heading("Contact Details", 1)
doc.add_paragraph(f"📞 Phone: {phone}")
doc.add_paragraph(f"📧 Email: {email}")
doc.add_paragraph(f"🏠 Address: {address}")

doc.add_heading("Education", 1)
doc.add_paragraph(f"{education} - {college}")

doc.add_heading("Experience", 1)
doc.add_paragraph(experience)

doc.add_heading("Skills", 1)
doc.add_paragraph(skills)

doc.add_heading("Projects", 1)
doc.add_paragraph(projects)

doc.add_heading("Languages", 1)
doc.add_paragraph(languages)

doc.add_heading("Hobbies", 1)
doc.add_paragraph(hobbies)

# RATINGS SECTION - separate at the end
doc.add_heading("Ratings", 1)
doc.add_paragraph(f"🛠 Skills Rating:      {skill_stars}")
doc.add_paragraph(f"📁 Projects Rating:   {project_stars}")
doc.add_paragraph(f"🌐 Languages Rating: {language_stars}")
doc.add_paragraph(f"💼 Experience Rating: {experience_stars}")
doc.save("C:\\Users\\Hp\\Desktop\\Resume.docx")
print("Details added successfully")







